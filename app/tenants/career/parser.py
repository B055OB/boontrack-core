"""app/tenants/career/parser.py
In-Memory Resume & Document Parser for Career Pipeline.
Extracts clean text directly from memory streams (io.BytesIO) with zero disk I/O.
"""

import io
import re
import logging
from typing import Union
from pypdf import PdfReader

logger = logging.getLogger("CAREER_PARSER")

MAGIC_BYTES_PDF = b"%PDF"
MAGIC_BYTES_ZIP = b"PK\x03\x04"  # DOCX


def parse_resume_buffer(file_buffer: Union[bytes, io.BytesIO], filename: str = "") -> str:
    """
    Extracts text from PDF or DOCX buffer in-memory with zero disk write.
    Uses pypdf with layout mode fallback and docx paragraph/table extraction.
    """
    if isinstance(file_buffer, io.BytesIO):
        file_bytes = file_buffer.getvalue()
    elif isinstance(file_buffer, bytes):
        file_bytes = file_buffer
    elif hasattr(file_buffer, "read"):
        file_bytes = file_buffer.read()
        if isinstance(file_bytes, str):
            file_bytes = file_bytes.encode("utf-8")
    else:
        raise TypeError(f"Unsupported buffer type: {type(file_buffer)}")

    if not file_bytes or len(file_bytes) < 4:
        logger.warning(f"[Career Parser] Buffer is empty or too small ({len(file_bytes) if file_bytes else 0} bytes)")
        return ""

    filename_lower = str(filename or "").lower().strip()
    is_pdf = file_bytes.startswith(MAGIC_BYTES_PDF) or filename_lower.endswith(".pdf")
    is_docx = file_bytes.startswith(MAGIC_BYTES_ZIP) or filename_lower.endswith(".docx")
    extracted_text = ""

    # 1. In-Memory PDF Extraction
    if is_pdf:
        try:
            reader = PdfReader(io.BytesIO(file_bytes), strict=False)
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception as enc_err:
                    logger.warning(f"[Career Parser] Encrypted PDF, decrypt attempt: {enc_err}")

            page_texts = []
            for page in reader.pages:
                txt = page.extract_text() or ""
                if not txt.strip():
                    try:
                        txt = page.extract_text(extraction_mode="layout") or ""
                    except Exception:
                        pass
                if txt.strip():
                    page_texts.append(txt.strip())

            if page_texts:
                extracted_text = "\n\n".join(page_texts)
        except Exception as e:
            logger.error(f"[Career Parser] Failed extracting PDF {filename}: {e}")

    # 2. In-Memory DOCX Extraction
    elif is_docx:
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_bytes))
            para_texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            table_texts = []
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        table_texts.append(" | ".join(cells))
            all_parts = para_texts + table_texts
            if all_parts:
                extracted_text = "\n\n".join(all_parts)
        except Exception as e:
            logger.error(f"[Career Parser] Failed extracting DOCX {filename}: {e}")

    # 3. Plain Text Fallback
    if not extracted_text.strip():
        if not file_bytes.startswith(b"\x00") and not file_bytes.startswith(MAGIC_BYTES_PDF) and not file_bytes.startswith(MAGIC_BYTES_ZIP):
            for enc in ("utf-8", "latin-1"):
                try:
                    decoded = file_bytes.decode(enc).strip()
                    if len(decoded) > 20 and not re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", decoded):
                        extracted_text = decoded
                        break
                except Exception:
                    pass

    clean_text = extracted_text.strip()
    logger.info(f"[Career Parser] Extracted {len(clean_text)} chars from '{filename}'")
    return clean_text
